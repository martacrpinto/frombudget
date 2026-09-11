from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
TEMP = Path(r"C:\Users\mpinto\AppData\Local\Temp")

BLACK = "101010"
SAGE = "78917A"
CREAM = "F7F5F0"
INK = "242424"
MUTED = "6B706C"
LIGHT = "E7E8E4"

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_text(cell, text, bold=False, color=INK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_doc_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def make_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10)
    styles['Normal'].font.color.rgb = RGBColor.from_string(INK)
    for name, size, color in [('Title', 25, BLACK), ('Heading 1', 17, BLACK), ('Heading 2', 12.5, SAGE)]:
        styles[name].font.name = 'Aptos Display' if name != 'Heading 2' else 'Aptos'
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor.from_string(color)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Budget Solution Cloud  |  Transfer guide').font.size = Pt(8)

    p = doc.add_paragraph(style='Title')
    p.add_run('Transferir a Budget Solution')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run('Guia para passar o projeto para as contas Supabase e Vercel do novo proprietário')
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(SAGE)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    facts = [
        ('Aplicação atual', 'https://frombudget.vercel.app'),
        ('Repositório', 'https://github.com/martacrpinto/frombudget'),
        ('Frontend', 'Vite + React na pasta client'),
        ('Autenticação', 'Supabase Auth com credenciais entregues manualmente'),
    ]
    for i, (a, b) in enumerate(facts):
        set_cell_text(table.cell(i, 0), a, True, SAGE, 9.5)
        set_cell_text(table.cell(i, 1), b, False, INK, 9.5)
        shade_cell(table.cell(i, 0), 'F0F2EE')
    doc.add_paragraph()

    doc.add_heading('1. Antes de começar', level=1)
    doc.add_paragraph('Confirma que o novo proprietário tem acesso às três áreas: GitHub, Supabase e Vercel. Guarda as chaves num password manager e mantém o projeto atual ativo até a nova conta passar os testes.')
    add_doc_bullet(doc, 'Não enviar para GitHub ficheiros .env, bases SQLite, backups, node_modules ou builds.')
    add_doc_bullet(doc, 'Manter uma cópia intacta do backup SQLite e um export validado do Postgres.')
    add_doc_bullet(doc, 'Decidir se a transferência mantém o projeto Supabase atual ou se cria um projeto Supabase novo. Manter o projeto atual é o caminho mais simples e preserva os utilizadores.')

    doc.add_heading('2. Transferir GitHub', level=1)
    doc.add_paragraph('No GitHub, transfere o repositório privado para a conta ou organização do novo proprietário, ou adiciona-o como administrador e deixa a transferência formal para o fim.')
    add_doc_bullet(doc, 'Confirmar que o repositório continua privado.')
    add_doc_bullet(doc, 'Confirmar que a branch main é a branch de produção.')
    add_doc_bullet(doc, 'Confirmar que o último commit contém a opção administrativa Change email.')
    add_doc_bullet(doc, 'Não alterar o histórico nem apagar a branch main durante a transição.')

    doc.add_heading('3. Supabase: manter o projeto atual', level=1)
    doc.add_paragraph('Esta é a opção recomendada quando os dados e utilizadores atuais devem continuar a funcionar. O proprietário atual deve adicionar o novo proprietário à organização Supabase com a função adequada e, se a organização permitir, transferir a propriedade do projeto.')
    add_doc_bullet(doc, 'Confirmar acesso ao projeto Budget Solution e à branch main / produção.')
    add_doc_bullet(doc, 'Rever membros da organização, billing e permissões.')
    add_doc_bullet(doc, 'Rever as funções Edge, em especial admin-users, e os secrets da função.')
    add_doc_bullet(doc, 'Confirmar que o URL do projeto, Auth e Storage continuam iguais.')

    doc.add_heading('4. Supabase: criar um projeto novo', level=1)
    doc.add_paragraph('Usa esta opção apenas se o novo proprietário precisar de uma organização Supabase separada. O projeto novo começa vazio e requer migração e validação.')
    add_doc_bullet(doc, 'Criar um projeto Free novo e guardar o Project URL, anon key, secret/service-role key e ligação de base de dados em local seguro.')
    add_doc_bullet(doc, 'Aplicar as migrations pela ordem 001 a 005 da pasta supabase/migrations.')
    add_doc_bullet(doc, 'Importar uma cópia dos dados. Nunca mover ou alterar o backup original.')
    add_doc_bullet(doc, 'Criar os utilizadores no Auth, associando cada conta ao profile_id / legacy_id correto.')
    add_doc_bullet(doc, 'Criar os buckets privados avatars e exports e verificar as respetivas policies.')
    add_doc_bullet(doc, 'Ativar Realtime apenas para budget_entries, categories, page_comments e page_status.')
    add_doc_bullet(doc, 'Publicar a Edge Function admin-users no projeto novo.')

    doc.add_heading('5. Secrets e Edge Function', level=1)
    doc.add_paragraph('A função admin-users faz operações privilegiadas, como criar logins, gerar passwords temporárias e alterar emails. A chave secreta nunca entra no frontend.')
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(['Variável', 'Onde fica', 'Regra']):
        set_cell_text(t.cell(0, i), h, True, 'FFFFFF', 9); shade_cell(t.cell(0, i), BLACK)
    rows = [
        ('VITE_SUPABASE_URL', 'Vercel: Production / Preview / Development', 'Pode ser usada pelo frontend'),
        ('VITE_SUPABASE_ANON_KEY', 'Vercel: Production / Preview / Development', 'Chave pública, sujeita a RLS'),
        ('SUPABASE_URL', 'Secrets da Edge Function', 'Não colocar em VITE_*'),
        ('SUPABASE_SECRET_KEY ou SUPABASE_SERVICE_ROLE_KEY', 'Secrets da Edge Function', 'Nunca expor no browser ou GitHub'),
        ('APP_URL', 'Secrets da Edge Function', 'URL Vercel usada pelo CORS, se configurada'),
    ]
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row): set_cell_text(cells[i], val, False, INK, 8.5)
    doc.add_paragraph()
    doc.add_paragraph('Depois do deploy, testa na função: list, criar login para um perfil existente, set-password, change-email, update de papéis e delete protegido.')

    doc.add_heading('6. Transferir Vercel', level=1)
    add_doc_bullet(doc, 'Importar o repositório privado para a conta Vercel do novo proprietário ou transferir o projeto atual.')
    add_doc_bullet(doc, 'Usar client como Root Directory.')
    add_doc_bullet(doc, 'Manter o build Vite e o fallback SPA definido em client/vercel.json.')
    add_doc_bullet(doc, 'Adicionar apenas VITE_SUPABASE_URL e VITE_SUPABASE_ANON_KEY ao frontend.')
    add_doc_bullet(doc, 'Configurar Production, Preview e Development separadamente.')
    add_doc_bullet(doc, 'Atualizar no Supabase Auth os URLs autorizados e redirects para o novo domínio, se o domínio mudar.')
    add_doc_bullet(doc, 'Publicar primeiro um Preview e só depois promover main para produção.')

    doc.add_heading('7. Utilizadores e credenciais', level=1)
    doc.add_paragraph('O projeto usa um fluxo manual. Não é necessário enviar convites ou emails automáticos.')
    add_doc_bullet(doc, 'Em Manage Users, usa @ / Set up login para um perfil sem login.')
    add_doc_bullet(doc, 'Entrega à pessoa o email e a password temporária mostrados uma única vez.')
    add_doc_bullet(doc, 'No primeiro acesso a aplicação obriga a escolher uma password pessoal.')
    add_doc_bullet(doc, 'Para repor a password, usa o botão de cadeado e entrega a nova credencial temporária.')
    add_doc_bullet(doc, 'Para mudar o email, usa @ / Change email. A password mantém-se e não há envio de email.')

    doc.add_heading('8. Validação antes da troca', level=1)
    checks = [
        'Acesso anónimo bloqueado; login, logout e refresh funcionam.',
        '7 utilizadores, 12 categorias, 2 anos e 926 entradas coincidem.',
        'Auditoria, comentários, estados, papéis, permissions e FTEs coincidem com a cópia validada.',
        'Totais por ano, utilizador, categoria e mês coincidem.',
        'Autosave, colagem Excel, dashboards, remuneração e exports XLSX funcionam.',
        'Dois browsers recebem atualizações Realtime.',
        'Admin, approver e utilizador normal veem apenas o que lhes compete.',
        'Alterar email permite login com o novo email e a password anterior.',
        'O bundle Vercel não contém secret key, service role key ou ligação de base de dados.',
    ]
    for c in checks: add_doc_bullet(doc, c)

    doc.add_heading('9. Troca e rollback', level=1)
    doc.add_paragraph('Mantém o projeto antigo disponível durante o piloto. Depois da aprovação dos administradores, atualiza o link usado pelos utilizadores e entrega as credenciais temporárias. Se algo divergir, bloqueia novas escritas, preserva logs e snapshots e volta ao deployment anterior. Não apagues imediatamente o projeto Supabase antigo.')

    doc.add_heading('10. Entrega ao novo proprietário', level=1)
    t2 = doc.add_table(rows=1, cols=2); t2.style = 'Table Grid'; t2.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(['Item', 'Confirmado por']): set_cell_text(t2.cell(0, i), h, True, 'FFFFFF', 9); shade_cell(t2.cell(0, i), BLACK)
    for item in ['GitHub privado e branch main', 'Projeto Supabase, membros e secrets', 'Edge Function admin-users', 'Projeto Vercel e variáveis', 'Backup SQLite e export Postgres', 'Lista nome-email dos utilizadores', 'URL final e procedimento de rollback']:
        cells = t2.add_row().cells; set_cell_text(cells[0], item, False, INK, 9); set_cell_text(cells[1], '[nome / data]', False, MUTED, 9)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run('Nota de segurança: '); r.bold = True; r.font.color.rgb = RGBColor.from_string(SAGE)
    p.add_run('as passwords temporárias devem ser entregues por um canal privado e não devem ser guardadas em documentos, GitHub ou screenshots.')

    path = OUT / 'Budget_Solution_Transfer_Guide.docx'
    doc.save(path)
    return path

def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False, font='Aptos', align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = PInches(0.03); tf.margin_right = PInches(0.03)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = font; r.font.size = PPt(size); r.font.bold = bold; r.font.color.rgb = PRGB.from_string(color)
    return box

def add_title(slide, title, kicker=None):
    if kicker: add_text(slide, kicker.upper(), 0.65, 0.35, 11.5, 0.25, 9, SAGE, True)
    add_text(slide, title, 0.65, 0.65 if kicker else 0.45, 11.7, 0.62, 28, BLACK, True, 'Aptos Display')
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PInches(0.65), PInches(1.42 if kicker else 1.22), PInches(1.0), PInches(0.05))
    line.fill.solid(); line.fill.fore_color.rgb = PRGB.from_string(SAGE); line.line.fill.background()

def add_footer(slide, n):
    add_text(slide, 'from: · Budget Solution', 0.65, 7.05, 5, 0.2, 8, MUTED)
    add_text(slide, f'{n:02d}', 12.0, 7.05, 0.6, 0.2, 8, MUTED, False, 'Aptos', PP_ALIGN.RIGHT)

def add_card(slide, x, y, w, h, title, body, accent=SAGE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PInches(x), PInches(y), PInches(w), PInches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = PRGB.from_string('FFFFFF'); shape.line.color.rgb = PRGB.from_string('D9DDD7')
    add_text(slide, title, x+0.22, y+0.2, w-0.44, 0.32, 15, accent, True)
    add_text(slide, body, x+0.22, y+0.66, w-0.44, h-0.8, 11.5, INK)

def make_ppt():
    prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]
    def base():
        s = prs.slides.add_slide(blank); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=PRGB.from_string(CREAM); return s
    # 1
    s=base(); add_text(s,'from:',0.75,1.15,7,1.0,54,BLACK,True,'Georgia'); add_text(s,'BUDGET SYSTEM',0.82,2.35,5,0.3,12,SAGE,True); add_text(s,'Manual de utilização',0.82,3.05,8,0.75,32,BLACK,True,'Aptos Display'); add_text(s,'Como entrar, preencher, validar e exportar budgets',0.84,3.95,8,0.35,15,MUTED); add_text(s,'https://frombudget.vercel.app',0.84,6.5,7,0.3,12,SAGE,True); add_footer(s,1)
    # 2
    s=base(); add_title(s,'Entrar na aplicação','01 · acesso'); add_card(s,0.7,1.9,5.7,3.5,'Acesso','Abra o endereço da aplicação.\n\nIntroduza o email e a password recebidos do administrador e clique em Enter.\n\nO sistema não envia convites automáticos. As credenciais são entregues manualmente.'); img=TEMP/'codex-clipboard-a9e13942-94a6-497c-89e6-d0482a5dc822.png';
    if img.exists(): s.shapes.add_picture(str(img),PInches(7.0),PInches(1.65),width=PInches(4.9),height=PInches(4.65)); add_footer(s,2)
    # 3
    s=base(); add_title(s,'Primeiro acesso','02 · password'); add_text(s,'As credenciais temporárias servem apenas para entrar pela primeira vez.',0.75,1.8,11.4,0.4,17,INK); add_card(s,0.8,2.55,3.55,2.35,'1 · Entrar','Use o email e a password temporária entregues pelo administrador.'); add_card(s,4.9,2.55,3.55,2.35,'2 · Definir','Escolha uma password pessoal e confirme-a.'); add_card(s,9.0,2.55,3.55,2.35,'3 · Continuar','A partir daí, use sempre a password pessoal.'); add_text(s,'Se perder a password, peça ao administrador novas credenciais temporárias.',0.85,5.65,10.8,0.4,14,SAGE,True); add_footer(s,3)
    # 4
    s=base(); add_title(s,'Navegação principal','03 · estrutura'); add_text(s,'A app organiza o trabalho por ano, categoria e página de budget.',0.75,1.72,11,0.4,17,INK); add_card(s,0.8,2.45,3.55,2.8,'Ano','Escolha o ano do budget que pretende consultar ou preencher.'); add_card(s,4.9,2.45,3.55,2.8,'Categoria','Abra a categoria e a página correspondente.'); add_card(s,9.0,2.45,3.55,2.8,'Totais','Use dashboards para acompanhar From, Heads e comparações.'); add_footer(s,4)
    # 5
    s=base(); add_title(s,'Preencher o budget','04 · edição'); add_card(s,0.8,1.85,5.5,3.5,'Valores mensais','Clique na célula pretendida e introduza o valor.\n\nA aplicação guarda alterações automaticamente. Confirme os totais depois de editar.\n\nOs valores ficam ligados ao ano, categoria, utilizador e mês.'); add_card(s,7.0,1.85,5.4,3.5,'Colar do Excel','Copie as células no Excel.\n\nSelecione a primeira célula correspondente na app e cole.\n\nReveja linhas, colunas e totais após a colagem.'); add_footer(s,5)
    # 6
    s=base(); add_title(s,'Notas, comentários e estado','05 · colaboração'); add_card(s,0.8,2.0,3.55,2.85,'Notas de linha','Explique um valor diretamente na linha quando for preciso deixar contexto.'); add_card(s,4.9,2.0,3.55,2.85,'Comentários','Use os comentários da página para perguntas, respostas e esclarecimentos.'); add_card(s,9.0,2.0,3.55,2.85,'Estado','Atualize o progresso da página e acompanhe alterações dos restantes utilizadores.'); add_footer(s,6)
    # 7
    s=base(); add_title(s,'Submeter e validar','06 · controlo'); add_text(s,'Antes de submeter, reveja valores, notas e comentários.',0.8,1.75,11.2,0.4,17,INK); add_card(s,0.85,2.55,3.55,2.55,'Submeter','Envie a página para validação quando os dados estiverem prontos.'); add_card(s,4.9,2.55,3.55,2.55,'Validar','Um approver ou administrador confirma a página ou pede correções.'); add_card(s,8.95,2.55,3.55,2.55,'Bloquear','Depois de validada, a página fica bloqueada para edição.'); add_footer(s,7)
    # 8
    s=base(); add_title(s,'Dashboards e exportação','07 · resultados'); add_card(s,0.8,1.9,5.5,3.35,'Acompanhar','Use os dashboards para comparar totais, anos e áreas. O simulador de remuneração permite testar configurações e FTEs.'); add_card(s,7.0,1.9,5.4,3.35,'Exportar XLSX','Abra a opção de exportação, escolha o âmbito e gere o ficheiro Excel. Guarde a versão para consulta ou partilha.'); add_footer(s,8)
    # 9
    s=base(); add_title(s,'Tarefas do administrador','08 · gestão'); add_card(s,0.8,1.75,5.5,3.7,'Credenciais manuais','Em Manage Users, use @ para criar um login num perfil sem email. Use o cadeado para gerar uma nova password temporária. Entregue as credenciais manualmente.'); add_card(s,7.0,1.75,5.4,3.7,'Alterar email','Use @ / Change email. A password mantém-se igual e não é enviado nenhum email. A pessoa deve entrar depois com o novo endereço.'); add_footer(s,9)
    # 10
    s=base(); add_title(s,'Se algo não funcionar','09 · ajuda'); add_card(s,0.8,1.85,5.5,3.7,'Primeiras verificações','Atualize a página. Confirme o ano e a categoria. Se os totais não atualizarem, aguarde alguns segundos e volte a abrir a página.'); add_card(s,7.0,1.85,5.4,3.7,'Falar com o administrador','Peça novas credenciais temporárias se não conseguir entrar. Para problemas de dados, indique o ano, a categoria e a linha afetada.'); add_text(s,'URL da aplicação  ·  https://frombudget.vercel.app',0.85,6.25,11,0.35,14,SAGE,True); add_footer(s,10)
    path=OUT/'Budget_Solution_User_Manual.pptx'; prs.save(path); return path

if __name__ == '__main__':
    print(make_ppt())
    print(make_doc())
